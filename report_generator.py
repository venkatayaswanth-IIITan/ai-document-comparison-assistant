"""
report_generator.py
-------------------
Generates a downloadable, professionally formatted DOCX report
using python-docx based on the comparison results.
"""

from datetime import datetime
from io import BytesIO
from typing import Dict, Any, Optional

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn


# Color Palette for Word Document
COLOR_PRIMARY_HEX = "1E3A8A"      # Deep Navy Blue
COLOR_SECONDARY_HEX = "2563EB"    # Royal Blue
COLOR_ACCENT_HEX = "0D9488"       # Teal
COLOR_BG_LIGHT_HEX = "F1F5F9"     # Light Slate / Gray
COLOR_MUTED_HEX = "64748B"        # Muted Gray

COLOR_PRIMARY = RGBColor(30, 58, 138)
COLOR_SECONDARY = RGBColor(37, 99, 235)
COLOR_TEXT_MAIN = RGBColor(15, 23, 42)
COLOR_TEXT_MUTED = RGBColor(100, 116, 139)


def set_cell_background(cell, hex_color: str):
    """Sets background shading color for a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)


def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
    """Sets cell padding (in dxa: 20 dxa = 1 pt)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for margin_name, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(margin_name)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)


def add_styled_heading(doc, text: str, level: int):
    """Adds a custom styled heading with theme color."""
    p = doc.add_heading(text, level=level)
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    for run in p.runs:
        if level == 1:
            run.font.color.rgb = COLOR_PRIMARY
            run.font.size = Pt(16)
            run.font.bold = True
        elif level == 2:
            run.font.color.rgb = COLOR_SECONDARY
            run.font.size = Pt(13)
            run.font.bold = True
        else:
            run.font.color.rgb = COLOR_TEXT_MAIN
            run.font.size = Pt(11)
            run.font.bold = True
    return p


def render_markdown_section_to_docx(doc, section_text: str):
    """
    Translates basic markdown formatting (bullet points, sub-headings, bold text)
    into native Word paragraphs and bullet lists.
    """
    if not section_text or not section_text.strip():
        p = doc.add_paragraph("No data provided for this section.")
        p.runs[0].font.italic = True
        p.runs[0].font.color.rgb = COLOR_TEXT_MUTED
        return

    lines = section_text.split("\n")
    for raw_line in lines:
        line = raw_line.strip()
        if not line:
            continue

        # Handle subheadings within markdown
        if line.startswith("### "):
            add_styled_heading(doc, line[4:].strip(), level=3)
        elif line.startswith("## "):
            add_styled_heading(doc, line[3:].strip(), level=2)
        elif line.startswith("# "):
            add_styled_heading(doc, line[2:].strip(), level=1)
        elif line.startswith("- ") or line.startswith("* ") or line.startswith("• "):
            # Bullet point
            clean_text = line[2:].strip()
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            _add_formatted_runs(p, clean_text)
        elif line[0].isdigit() and len(line) > 2 and line[1] in [".", ")"]:
            # Numbered item
            clean_text = line[2:].strip()
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            _add_formatted_runs(p, clean_text)
        else:
            # Regular body paragraph
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(5)
            p.paragraph_format.line_spacing = 1.15
            _add_formatted_runs(p, line)


def _add_formatted_runs(paragraph, text: str):
    """Parses simple inline bold **text** and creates styled runs."""
    import re
    tokens = re.split(r'(\*\*.*?\*\*)', text)
    for token in tokens:
        if token.startswith('**') and token.endswith('**') and len(token) > 4:
            run = paragraph.add_run(token[2:-2])
            run.bold = True
            run.font.color.rgb = COLOR_TEXT_MAIN
        else:
            run = paragraph.add_run(token)
            run.font.color.rgb = COLOR_TEXT_MAIN


def create_comparison_docx(
    doc_a_meta: Dict[str, Any],
    doc_b_meta: Dict[str, Any],
    comparison_data: Dict[str, str],
    model_name: str = "gemini-2.5-flash"
) -> BytesIO:
    """
    Builds a professional DOCX document containing the full document comparison report.

    Args:
        doc_a_meta: Dict with filename, page_count, char_count, word_count for Doc A.
        doc_b_meta: Dict with filename, page_count, char_count, word_count for Doc B.
        comparison_data: Dict containing summary_a, summary_b, similarities,
                         differences, key_details, overall_conclusion, limitations.
        model_name: Name of AI model used for generation.

    Returns:
        BytesIO object containing the complete DOCX file.
    """
    doc = docx.Document()

    # Set page margins (1 inch)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

        # Header / Footer
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("AI Document Comparison Assistant | Confidential Report")
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = COLOR_TEXT_MUTED

        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run(f"Generated on {datetime.now().strftime('%Y-%m-%d %H:%M:%S')} • Grounded in Extracted Text Evidence")
        frun.font.size = Pt(8.5)
        frun.font.color.rgb = COLOR_TEXT_MUTED

    # Document Header Title
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(2)
    title_run = title_p.add_run("AI Document Comparison Assistant")
    title_run.font.size = Pt(24)
    title_run.font.bold = True
    title_run.font.color.rgb = COLOR_PRIMARY

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_before = Pt(0)
    sub_p.paragraph_format.space_after = Pt(14)
    sub_run = sub_p.add_run("Automated Comparative Analysis, Differences, and Evidence Matrix")
    sub_run.font.size = Pt(12)
    sub_run.font.color.rgb = COLOR_TEXT_MUTED

    # Metadata Summary Box / Table
    add_styled_heading(doc, "Document Metadata & Run Configuration", level=2)
    meta_table = doc.add_table(rows=5, cols=3)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False

    # Widths
    col_widths = [Inches(2.2), Inches(2.1), Inches(2.1)]

    headers = ["Parameter / Metric", "Document A (Baseline)", "Document B (Comparison)"]
    hdr_cells = meta_table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_background(hdr_cells[i], COLOR_PRIMARY_HEX)
        set_cell_margins(hdr_cells[i])
        for p in hdr_cells[i].paragraphs:
            for r in p.runs:
                r.font.bold = True
                r.font.color.rgb = RGBColor(255, 255, 255)
                r.font.size = Pt(9.5)

    rows_data = [
        ("Filename", doc_a_meta.get("filename", "N/A"), doc_b_meta.get("filename", "N/A")),
        ("Page Count", f"{doc_a_meta.get('page_count', 0)} page(s)", f"{doc_b_meta.get('page_count', 0)} page(s)"),
        ("Character Count", f"{doc_a_meta.get('char_count', 0):,} chars", f"{doc_b_meta.get('char_count', 0):,} chars"),
        ("Word Count", f"{doc_a_meta.get('word_count', 0):,} words", f"{doc_b_meta.get('word_count', 0):,} words"),
    ]

    for row_idx, row_content in enumerate(rows_data, start=1):
        row_cells = meta_table.rows[row_idx].cells
        for col_idx, text in enumerate(row_content):
            row_cells[col_idx].text = str(text)
            set_cell_margins(row_cells[col_idx])
            # Zebra striping
            if row_idx % 2 == 1:
                set_cell_background(row_cells[col_idx], COLOR_BG_LIGHT_HEX)
            for p in row_cells[col_idx].paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9)
                    if col_idx == 0:
                        r.font.bold = True
                        r.font.color.rgb = COLOR_TEXT_MAIN

    # Additional execution details paragraph
    detail_p = doc.add_paragraph()
    detail_p.paragraph_format.space_before = Pt(8)
    detail_p.paragraph_format.space_after = Pt(16)
    d_run = detail_p.add_run(f"AI Model: {model_name} | Generated At: {datetime.now().strftime('%B %d, %Y at %I:%M %p')}")
    d_run.font.size = Pt(8.5)
    d_run.font.italic = True
    d_run.font.color.rgb = COLOR_TEXT_MUTED

    # 1. Summary of Document A
    add_styled_heading(doc, f"1. Executive Summary: Document A ({doc_a_meta.get('filename', 'Baseline')})", level=1)
    render_markdown_section_to_docx(doc, comparison_data.get("summary_a", ""))

    # 2. Summary of Document B
    add_styled_heading(doc, f"2. Executive Summary: Document B ({doc_b_meta.get('filename', 'Comparison')})", level=1)
    render_markdown_section_to_docx(doc, comparison_data.get("summary_b", ""))

    # 3. Similarities
    add_styled_heading(doc, "3. Shared Provisions & Key Similarities", level=1)
    render_markdown_section_to_docx(doc, comparison_data.get("similarities", ""))

    # 4. Differences & Contradictions
    add_styled_heading(doc, "4. Differences, Additions, & Contradictions", level=1)
    render_markdown_section_to_docx(doc, comparison_data.get("differences", ""))

    # 5. Key Details Comparison
    add_styled_heading(doc, "5. Key Details Matrix (Dates, Names, Amounts, Deadlines)", level=1)
    render_markdown_section_to_docx(doc, comparison_data.get("key_details", ""))

    # 6. Overall Conclusion
    add_styled_heading(doc, "6. Overall Conclusion & Synthesis", level=1)
    render_markdown_section_to_docx(doc, comparison_data.get("overall_conclusion", ""))

    # 7. Limitations & Grounding
    add_styled_heading(doc, "7. Limitations & Grounding Statement", level=1)
    render_markdown_section_to_docx(doc, comparison_data.get("limitations", ""))

    # Save to BytesIO stream
    output_stream = BytesIO()
    doc.save(output_stream)
    output_stream.seek(0)
    return output_stream
